# -*- coding: utf-8 -*-
"""
페피트 메뉴판 빌더
------------------
data/*.json 의 내용으로 docx 메뉴판을 생성한다.

레이아웃 원칙
  - 메뉴 1개 = 표 1행 (왼쪽 셀: 이름/설명, 오른쪽 셀: 가격)
  - 표 테두리는 전부 없음 -> 화면상 보이지 않지만 위치는 고정
  - 행 분할 금지(cantSplit) -> 페이지 경계에서 한 메뉴가 잘리지 않음
  - 섹션 헤더는 keepNext -> 헤더만 페이지 끝에 남지 않음
  - 페이지 구성은 PAGES 에 명시 -> 앞 페이지 내용이 늘어도 뒤 페이지가 안 밀림

수정 방법
  - 메뉴 내용: data/*.json 편집
  - 글자 크기/여백/페이지 구성: 아래 STYLE, PAGES 편집
  - 실행:  python build.py           (docx 생성)
          python build.py --debug   (표 테두리를 보이게 해서 배치 확인)
"""
import json, os, sys

import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import (WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING,
                            WD_TAB_ALIGNMENT)
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, 'data')
DIST = os.path.join(HERE, 'dist')
os.makedirs(DIST, exist_ok=True)
DEBUG = '--debug' in sys.argv

# --------------------------------------------------------------------------
# 스타일 (여기만 고치면 전체 메뉴판에 반영됨)
# --------------------------------------------------------------------------
# 폰트는 원본 4개 파일이 문서마다 달랐다.
#   요리/런치/코스 = 바탕 + Cambria(테마 기본), 와인 = 맑은 고딕 + Times New Roman,
#   코스 영문에는 서울남산체 L 이 20군데 섞여 있었음.
# 다수(요리/런치/코스)의 명조 기조를 기준으로 전 문서를 통일했다.
STYLE = {
    'font_head': 'ChunkFive Ex',    # 로고 / 섹션 헤더 / 가격
    'font_kor': '바탕',              # 한글 본문 (원본 요리/런치/코스와 동일)
    'font_eng': 'Times New Roman',  # 영문 본문 (원본 와인과 동일, Cambria 와 같은 명조 계열)

    'title_size': 22,              # PÉPITE
    'section_size': 18,            # STARTER, PASTA ...
    'name_size': 11.5,             # 메뉴 한글명
    'engname_size': 9,             # 메뉴 영문명
    'desc_size': 9,                # 설명
    'price_size': 11.5,
    'small_size': 7,               # 부가 표기 (kor_small, price_note, 코스의 작은 괄호)

    # 브랜드 색 (기존 메뉴판에서 그대로 가져옴)
    'logo': RGBColor(0xE3, 0x6C, 0x0A),    # PÉPITE 로고 주황
    'gray': RGBColor(0x40, 0x40, 0x40),    # 영문 보조 텍스트
    'header': RGBColor(0x00, 0x00, 0x00),  # 섹션 헤더는 검정

    # 제목 아래 가로선 (원본과 동일: single 0.5pt, 글자와 간격 14)
    'rule_under_title': True,
    'rule_under_section': False,   # 원본에는 없음. True 로 켜면 섹션 헤더 아래에도 선
    'rule_size': 4,                # 1/8 pt 단위 -> 4 = 0.5pt
    'rule_space': 14,

    'item_gap': 9,                 # 메뉴 사이 간격(pt)
    'section_gap_before': 16,      # 섹션 헤더 위 간격
    'section_gap_after': 8,

    'page_w': 21.0, 'page_h': 29.7,          # A4 (cm)
    'margin_lr': 2.2, 'margin_tb': 1.6,
    'price_col': 2.4,                        # 가격 칸 폭(cm)
}

# 와인은 항목 수가 많아 본문보다 한 단계 작게 (원본 와인 메뉴판과 비슷한 밀도)
# 한글명은 설명과 같은 크기로 두고 굵기로만 구분한다 (2026-07-29 요청).
WINE_STYLE = {
    'name_size': 8,     # = desc_size. 크기가 아니라 bold 로 제목임을 표시
    'engname_size': 8,
    'region_size': 7,
    'desc_size': 8,
    'price_size': 9.5,
    'item_gap': 16,     # 기본값. 실제 간격은 PAGES['wine'] 에서 페이지마다 지정한다.
}

# 위스키 메뉴판 (2026-08-17 수제 docx 에서 마이그레이션. 병 사진/3단 구성이 있어
# PAGES 를 쓰지 않고 전용 빌더 build_whisky 로 만든다. 1페이지 고정.)
# 폰트는 원본 그대로 Adobe 명조 Std M 을 유지한다. 일본어 신자체 한자(焼 등)가
# 있어서 바탕으로 바꾸면 글자가 깨질 수 있고, 미설치 머신에서는 원본 docx 때와
# 동일하게 Word 가 대체 폰트로 그린다.
WHISKY_STYLE = {
    'font': 'Adobe 명조 Std M',
    'header_size': 20,        # 섹션 헤더 (日本焼酎 JAPANESE SHOCHU / WHISKY / ETC.)
    'category_size': 16,      # 고구마 / 보리 / 쌀
    'name_size': 12,          # 품목명
    'sub_size': 8,            # 한자 표기 + 영문명
    'spec_size': 7,           # 알코올도수 / 원산지 / 용량
    'desc_size': 9,           # 위스키 설명문
    'label_size': 8,          # BOTTLE, Shot / Highball
    'price_size': 10.5,       # 위스키 가격
    'shochu_price_size': 11,  # 쇼츄 Bottle 가격
    'etc_size': 12,           # ETC 품목 (원본 14 + 장평 83%. 장평 없이 12 로 비슷한 폭)
    'margin_lr': 2.2, 'margin_tb': 1.6,
}

# 페이지 구성. 리스트 하나 = 한 페이지.
#
#   페이지    ['STARTER', 'PASTA']                     섹션 목록
#             {'sections': [...], 'item_gap': 36}      그 페이지만 항목 간격 지정
#   섹션      'STARTER'                                섹션 전체
#             ('WHITE WINE', 0, 10)                    items[0:10] 만
#             ('WHITE WINE', 10, None, False)          items[10:], 헤더 없이
#
# item_gap 은 "페이지 하단에 여백을 남기지 않는다"는 원칙에 따라 페이지마다
# 정한 값이다. 항목 수가 적은 페이지는 간격이 넓다. 값을 바꾸면 반드시
# 빌드 후 PDF 로 페이지 수와 경계를 확인할 것.
PAGES = {
    'food': [
        {'sections': ['STARTER'], 'item_gap': 34},          # 7항목
        {'sections': ['PASTA'], 'item_gap': 53},            # 6항목
        {'sections': ['MAIN', 'MEAL'], 'item_gap': 28},     # 5+2항목
        {'sections': ['SIDE', 'DESSERT'], 'item_gap': 34},  # 5+2항목
    ],
    # 런치는 항목이 많아 2페이지에 넣으려면 기본 9pt 보다 좁혀야 한다.
    'lunch': [
        {'sections': ['STARTER', 'PASTA'], 'item_gap': 6},                     # 6+4항목
        {'sections': ['MAIN', 'MEAL', 'DESSERT', 'SIDE', 'DRINKS'],
         'item_gap': 6},                                                       # 3+1+2+1+2항목
    ],
    # 와인은 섹션당 항목이 많아 한 섹션이 한 페이지에 안 들어간다. 끊는 위치를
    # 여기에 명시해서 앞쪽에 한 종을 추가해도 뒤 페이지가 밀리지 않게 한다.
    # 2026-10-27 요청: 섹션은 페이지 중간에서 시작하지 않고 항상 새 페이지 맨
    # 위에서 시작한다 (섹션당 2페이지, 6페이지 유지). 그래서 화이트/레드 페이지는
    # item_gap 을 기본 16 보다 좁혔다.
    'wine': [
        {'sections': [('WHITE WINE', 0, 10)], 'item_gap': 10},
        {'sections': [('WHITE WINE', 10, None, False)], 'item_gap': 14},
        {'sections': [('RED WINE', 0, 9)], 'item_gap': 14},
        {'sections': [('RED WINE', 9, None, False)], 'item_gap': 14},
        {'sections': [('SPARKLING WINE', 0, 6)], 'item_gap': 50},
        {'sections': [('SPARKLING WINE', 6, None, False), 'NON-ALCOHOLIC', 'HOUSE WINE'],
         'item_gap': 36},
    ],
}

# 특정 섹션만 헤더 위 간격을 다르게 (기본값은 STYLE['section_gap_before'])
# 같은 페이지에 두 섹션이 붙어 있을 때 사이를 더 벌리는 용도.
SECTION_GAP_BEFORE = {
    ('food', 'MEAL'): 26,     # MAIN 과 MEAL 사이를 한 줄 더 띄운다
}

# 영문 설명 서식. 기본은 한글 설명과 같은 크기 + 회색 보조 텍스트.
# 여기에 적은 메뉴판은 영문 설명을 영문명과 같은 서식(engname_size, 검정)으로 맞춘다.
# 요리(디너)는 2026-07-29 요청으로 영문명과 크기/색을 통일했다.
DESC_ENG_AS_ENGNAME = {'food'}


def page_spec(page):
    """페이지 정의를 (섹션 목록, 항목 간격) 으로 정규화."""
    if isinstance(page, dict):
        return page['sections'], page.get('item_gap')
    return page, None


def section_spec(entry):
    """섹션 정의를 (이름, start, end, 헤더표시) 로 정규화."""
    if isinstance(entry, (list, tuple)):
        name, start, end = entry[0], entry[1], entry[2]
        return name, start, end, (entry[3] if len(entry) > 3 else True)
    return entry, 0, None, True


# --------------------------------------------------------------------------
# 저수준 헬퍼
# --------------------------------------------------------------------------
def set_cell_borders_none(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single' if DEBUG else 'none')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), 'FF0000' if DEBUG else 'auto')
        borders.append(e)
    tblPr.append(borders)


def set_cell_margins(table, top=0, bottom=0, left=0, right=0):
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for side, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement('w:' + side)
        e.set(qn('w:w'), str(int(val * 567)))   # cm -> twips
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    tblPr.append(mar)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))


def bottom_rule(p, color=None):
    """문단 아래 가로선(hr). 원본 메뉴판의 PÉPITE 밑줄과 동일한 속성."""
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), str(STYLE['rule_size']))
    b.set(qn('w:space'), str(STYLE['rule_space']))
    b.set(qn('w:color'), '%02X%02X%02X' % tuple(color or STYLE['logo']))
    bdr.append(b)
    pPr.append(bdr)
    return p


def keep_next(p, on=True):
    pPr = p._p.get_or_add_pPr()
    e = OxmlElement('w:keepNext')
    if not on:
        e.set(qn('w:val'), '0')
    pPr.append(e)
    pPr.append(OxmlElement('w:keepLines'))


def put_run(p, text, *, font, size, bold=False, color=None, eastasia=None):
    """문단에 run 하나 추가. 한 줄 안에서 글자 크기를 바꿀 때 나눠 쓴다."""
    r = p.add_run(text)
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    r._element.rPr.rFonts.set(qn('w:eastAsia'), eastasia or STYLE['font_kor'])
    return r


def line(cell, text, *, font, size, bold=False, color=None,
         after=0, first=False, align=None, small=None):
    """셀 안에 한 줄(문단) 추가. first=True 면 셀의 기본 빈 문단을 재사용.

    small 을 주면 같은 줄 뒤에 작은 글씨로 덧붙인다 (예: 빵추가 + (3pieces)).
    """
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    put_run(p, text, font=font, size=size, bold=bold, color=color)
    if small:
        put_run(p, ' ' + small, font=font, size=STYLE['small_size'],
                bold=bold, color=color)
    return p


def section_header(doc, name, gap_before=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(STYLE['section_gap_before'] if gap_before is None else gap_before)
    pf.space_after = Pt(STYLE['section_gap_after'])
    keep_next(p)
    r = p.add_run(name)
    r.font.name = STYLE['font_head']
    r.font.size = Pt(STYLE['section_size'])
    r.font.color.rgb = STYLE['header']
    r._element.rPr.rFonts.set(qn('w:eastAsia'), STYLE['font_head'])
    if STYLE['rule_under_section']:
        bottom_rule(p)
    return p


def doc_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    keep_next(p)
    r = p.add_run(text)
    r.font.name = STYLE['font_head']
    r.font.size = Pt(STYLE['title_size'])
    r.font.color.rgb = STYLE['logo']       # 브랜드 주황
    r._element.rPr.rFonts.set(qn('w:eastAsia'), STYLE['font_head'])
    if STYLE['rule_under_title']:
        bottom_rule(p)
    return p


def new_doc():
    doc = docx.Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(STYLE['page_w']), Cm(STYLE['page_h'])
    s.left_margin = s.right_margin = Cm(STYLE['margin_lr'])
    s.top_margin = s.bottom_margin = Cm(STYLE['margin_tb'])
    n = doc.styles['Normal']
    n.font.name = STYLE['font_kor']
    n.font.size = Pt(STYLE['desc_size'])
    n.element.rPr.rFonts.set(qn('w:eastAsia'), STYLE['font_kor'])
    return doc


def item_table(doc, n_rows):
    body_w = STYLE['page_w'] - 2 * STYLE['margin_lr']
    t = doc.add_table(rows=n_rows, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    set_cell_borders_none(t)
    set_cell_margins(t)
    for row in t.rows:
        no_split(row)
        row.cells[0].width = Cm(body_w - STYLE['price_col'])
        row.cells[1].width = Cm(STYLE['price_col'])
    return t


def page_break(doc):
    """페이지 나눔. 앞 페이지가 꽉 찬 경우 이 문단 자체가 다음 페이지로 밀려
    빈 페이지가 생기므로 높이를 1pt 로 고정한다."""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(1)
    r = p.add_run()
    r.font.size = Pt(1)
    r.add_break(WD_BREAK.PAGE)


def tail_paragraph(doc):
    """본문 끝에 1pt 문단을 명시한다. 문서가 표로 끝나면 Word 가 뒤에 기본 크기의
    빈 문단을 자동으로 붙이는데, 마지막 페이지가 꽉 찬 경우 그 문단만 다음 페이지로
    밀려 빈 페이지가 생긴다. (와인 6페이지에서 실제 발생, 2026-08-17)"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(1)
    p.add_run().font.size = Pt(1)


def price_cell(cell, price, size=None, note=None):
    """가격. note 를 주면 가격 밑에 작은 글씨로 한 줄 더 (예: 1piece 추가 6.5)."""
    line(cell, price, font=STYLE['font_head'], size=size or STYLE['price_size'],
         first=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    if note:
        line(cell, note, font=STYLE['font_eng'], size=STYLE['small_size'],
             color=STYLE['gray'], align=WD_ALIGN_PARAGRAPH.RIGHT)


# --------------------------------------------------------------------------
# 메뉴판별 조립
# --------------------------------------------------------------------------
def build_dish(key, out_name):
    data = json.load(open(os.path.join(DATA, key + '.json'), encoding='utf-8'))
    by_name = {s['name']: s for s in data['sections']}
    as_engname = key in DESC_ENG_AS_ENGNAME
    desc_eng_size = STYLE['engname_size'] if as_engname else STYLE['desc_size']
    desc_eng_color = None if as_engname else STYLE['gray']
    doc = new_doc()

    for pi, page in enumerate(PAGES[key]):
        if pi:
            page_break(doc)
        sections, gap = page_spec(page)
        gap = STYLE['item_gap'] if gap is None else gap
        doc_title(doc, data['title'])
        for entry in sections:
            sec, start, end, show_header = section_spec(entry)
            items = by_name[sec]['items'][start:end]
            if show_header:
                section_header(doc, sec, SECTION_GAP_BEFORE.get((key, sec)))
            t = item_table(doc, len(items))
            for row, it in zip(t.rows, items):
                c = row.cells[0]
                line(c, it['kor'], font=STYLE['font_kor'], size=STYLE['name_size'],
                     bold=True, first=True, after=1, small=it.get('kor_small'))
                if it['eng']:
                    line(c, it['eng'], font=STYLE['font_eng'],
                         size=STYLE['engname_size'], bold=True, after=2)
                if it['desc_kor']:
                    line(c, '* ' + it['desc_kor'], font=STYLE['font_kor'],
                         size=STYLE['desc_size'], after=1)
                if it['desc_eng']:
                    line(c, it['desc_eng'], font=STYLE['font_eng'],
                         size=desc_eng_size, color=desc_eng_color)
                c.paragraphs[-1].paragraph_format.space_after = Pt(gap)
                price_cell(row.cells[1], it['price'], note=it.get('price_note'))

    tail_paragraph(doc)
    doc.save(os.path.join(DIST, out_name))
    print('saved', out_name)


def build_wine(out_name):
    data = json.load(open(os.path.join(DATA, 'wine.json'), encoding='utf-8'))
    by_name = {s['name']: s for s in data['sections']}
    doc = new_doc()

    for pi, page in enumerate(PAGES['wine']):
        if pi:
            page_break(doc)
        sections, gap = page_spec(page)
        gap = WINE_STYLE['item_gap'] if gap is None else gap
        for entry in sections:
            sec, start, end, show_header = section_spec(entry)
            items = by_name[sec]['items'][start:end]
            if show_header:
                section_header(doc, sec, SECTION_GAP_BEFORE.get(('wine', sec)))
            t = item_table(doc, len(items))
            for row, it in zip(t.rows, items):
                c = row.cells[0]
                line(c, it['kor'], font=STYLE['font_kor'], size=WINE_STYLE['name_size'],
                     bold=True, first=True, after=0)
                if it['eng']:
                    line(c, it['eng'], font=STYLE['font_eng'],
                         size=WINE_STYLE['engname_size'], bold=True, after=1)
                if it.get('region'):
                    line(c, it['region'], font=STYLE['font_kor'],
                         size=WINE_STYLE['region_size'], color=STYLE['gray'], after=1)
                if it['desc_kor']:
                    line(c, it['desc_kor'], font=STYLE['font_kor'],
                         size=WINE_STYLE['desc_size'], after=0)
                if it['desc_eng']:
                    line(c, it['desc_eng'], font=STYLE['font_eng'],
                         size=WINE_STYLE['desc_size'], color=STYLE['gray'])
                c.paragraphs[-1].paragraph_format.space_after = Pt(gap)
                price_cell(row.cells[1], it['price'], WINE_STYLE['price_size'])

    tail_paragraph(doc)
    doc.save(os.path.join(DIST, out_name))
    print('saved', out_name)


def build_course(key, out_name, item_gap=0):
    """코스 메뉴 (A5, 가운데 정렬). item_gap 은 한글 메뉴명 줄 위 간격(pt)."""
    data = json.load(open(os.path.join(DATA, key + '.json'), encoding='utf-8'))
    doc = docx.Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(14.8), Cm(21.0)     # A5
    s.left_margin = s.right_margin = Cm(1.6)
    # 섹션 사이를 한 줄 더 띄우면서 마지막 줄이 2페이지로 넘어가 1.4 -> 1.2 로 줄였고,
    # 식전빵/타르타르 사이 'and' 한 줄이 늘어나 1.2 -> 1.0 으로 더 줄였다 (1페이지 유지)
    s.top_margin = s.bottom_margin = Cm(1.0)
    n = doc.styles['Normal']
    n.font.name = STYLE['font_kor']
    n.element.rPr.rFonts.set(qn('w:eastAsia'), STYLE['font_kor'])

    heads = {'STARTER', 'PASTA', 'MAIN', 'DESSERT'}

    def put(text, size, *, font=None, bold=False, before=0, after=0, color=None,
            small=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = 1.0
        keep_next(p)
        f = font or STYLE['font_kor']
        # 한글 대체 폰트는 헤드 폰트를 쓰는 줄(로고/섹션명)만 예외로 둔다
        ea = font if font == STYLE['font_head'] else STYLE['font_kor']
        put_run(p, text, font=f, size=size, bold=bold, color=color, eastasia=ea)
        if small:
            put_run(p, ' ' + small, font=f, size=STYLE['small_size'],
                    bold=bold, color=color, eastasia=ea)
        return p

    # PÉPITE / LUNCH SET MENU 를 한 덩어리로 보고 가로선은 그 아래에 둔다
    put(data['title'], 20, font=STYLE['font_head'], after=2, color=STYLE['logo'])
    sub = put(data['subtitle'], 12, font=STYLE['font_head'], after=6,
              color=STYLE['logo'])
    if STYLE['rule_under_title']:
        bottom_rule(sub)

    for ln in data['lines']:
        # ["식전 빵", "(빵 추가 2pcs 3.)"] 처럼 리스트면 뒷부분을 작은 글씨로 붙인다
        small = None
        if isinstance(ln, list):
            ln, small = ln[0], ln[1]
        if ln in heads:
            # 코스 섹션 사이는 한 줄 더 띄운다 (2026-07-29 요청). 12 -> 26
            put(ln, 12, font=STYLE['font_head'], before=26, after=6)
        elif ln in ('or', 'and'):
            # 'and' 는 식전빵과 비프타르타르 사이 간격/연결 표시 (2026-08-17 요청)
            put(ln, 8, color=STYLE['gray'], before=1, after=1)
        elif ln.startswith('98,000'):
            put(ln, 16, font=STYLE['font_head'], before=20, after=2)
        elif ln.startswith('(1인기준'):
            put(ln, 9, color=STYLE['gray'])
        elif ln.startswith('(') or ln.startswith('Pre-meal') or ln.startswith('(Add'):
            put(ln, 6, color=STYLE['gray'], after=1)
        elif ln[0].isascii():
            put(ln, 8, font=STYLE['font_eng'], color=STYLE['gray'], after=3)
        else:
            put(ln, 10, bold=True, before=item_gap, after=1, small=small)

    doc.save(os.path.join(DIST, out_name))
    print('saved', out_name)


def build_whisky(out_name):
    """위스키 메뉴판. 원본 수제 docx 의 배치(쇼츄 3단 / 위스키 목록 / ETC 2단)를
    표 기반으로 재현한다. 가격은 탭 개수가 아니라 오른쪽 탭 스톱으로 정렬한다."""
    data = json.load(open(os.path.join(DATA, 'whisky.json'), encoding='utf-8'))
    ws = WHISKY_STYLE
    assets = os.path.join(HERE, 'assets', 'whisky')
    body_w = STYLE['page_w'] - 2 * ws['margin_lr']

    doc = docx.Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(STYLE['page_w']), Cm(STYLE['page_h'])
    s.left_margin = s.right_margin = Cm(ws['margin_lr'])
    s.top_margin = s.bottom_margin = Cm(ws['margin_tb'])
    n = doc.styles['Normal']
    n.font.name = ws['font']
    n.font.size = Pt(ws['desc_size'])
    n.element.rPr.rFonts.set(qn('w:eastAsia'), ws['font'])

    def run(p, text, size, bold=False, color=None):
        return put_run(p, text, font=ws['font'], size=size, bold=bold,
                       color=color, eastasia=ws['font'])

    def para(cell, *, first=False, after=0, align=None, right_tab=None):
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(after)
        pf.line_spacing = 1.0
        if align is not None:
            p.alignment = align
        if right_tab is not None:
            pf.tab_stops.add_tab_stop(Cm(right_tab), WD_TAB_ALIGNMENT.RIGHT)
        return p

    def header(text, before=18):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.space_before = Pt(before)
        pf.space_after = Pt(10)
        keep_next(p)
        run(p, text, ws['header_size'], bold=True)
        bottom_rule(p, color=RGBColor(0, 0, 0))
        return p

    def left_bar(p):
        """카테고리(고구마/보리/쌀) 왼쪽의 굵은 세로바. 원본의 6pt 세로선 재현."""
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement('w:pBdr')
        b = OxmlElement('w:left')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '48')          # 1/8pt 단위 -> 6pt
        b.set(qn('w:space'), '4')
        b.set(qn('w:color'), '000000')
        bdr.append(b)
        pPr.append(bdr)

    def cell_bottom_border(cell):
        """위스키 품목명 아래의 가로선 (이름 칸 + 가격 칸에 걸쳐 이어진다)."""
        tcPr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement('w:tcBorders')
        b = OxmlElement('w:bottom')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), '000000')
        borders.append(b)
        tcPr.append(borders)

    def picture(cell, name, width_cm, height_cm=None, center=False):
        # 원본이 사진을 특정 크기로 늘려 쓰고 있어서 폭/높이를 모두 강제한다
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(os.path.join(assets, name), width=Cm(width_cm),
                                height=Cm(height_cm) if height_cm else None)

    def base_table(rows, cols):
        t = doc.add_table(rows=rows, cols=cols)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        t.autofit = False
        set_cell_borders_none(t)
        set_cell_margins(t)
        for r in t.rows:
            no_split(r)
        return t

    for sec in data['sections']:
        header(sec['name'], before=0 if sec is data['sections'][0] else 18)

        if sec['type'] == 'shochu':
            col_w = body_w / 3
            t = base_table(1, 3)
            for cell, it in zip(t.rows[0].cells, sec['items']):
                cell.width = Cm(col_w)
                p = para(cell, first=True, after=6)
                left_bar(p)
                run(p, it['category'], ws['category_size'], bold=True)
                nt = cell.add_table(rows=1, cols=2)
                nt.autofit = False
                set_cell_borders_none(nt)
                set_cell_margins(nt)
                ic, tc = nt.rows[0].cells
                ic.width = Cm(it['img_w'] + 0.35)
                tc.width = Cm(col_w - it['img_w'] - 0.35)
                picture(ic, it['img'], it['img_w'], it['img_h'])
                run(para(tc, first=True, after=2), it['kor'], ws['name_size'], bold=True)
                run(para(tc, after=2), it['sub'], ws['sub_size'])
                run(para(tc, after=4), it['spec'], ws['spec_size'])
                pp = para(tc, right_tab=col_w - it['img_w'] - 0.8)
                run(pp, it['price_label'] + '\t', ws['shochu_price_size'], bold=True)
                run(pp, it['price'], ws['shochu_price_size'], bold=True)

        elif sec['type'] == 'whisky':
            img_col, price_col = 1.7, 4.0
            for i, it in enumerate(sec['items']):
                t = base_table(2, 3)
                for r in t.rows:
                    r.cells[0].width = Cm(img_col)
                    r.cells[1].width = Cm(body_w - img_col - price_col)
                    r.cells[2].width = Cm(price_col)
                ic = t.cell(0, 0).merge(t.cell(1, 0))
                ic.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                picture(ic, it['img'], it['img_w'], it['img_h'], center=True)

                run(para(t.cell(0, 1), first=True, after=1), it['kor'],
                    ws['name_size'], bold=True)
                pp = para(t.cell(0, 2), first=True, after=1,
                          right_tab=price_col - 0.15)
                run(pp, 'BOTTLE\t', ws['label_size'])
                run(pp, it['bottle'], ws['price_size'])
                cell_bottom_border(t.cell(0, 1))
                cell_bottom_border(t.cell(0, 2))

                body = t.cell(1, 1)
                run(para(body, first=True, after=1), it['eng'], ws['sub_size'])
                run(para(body, after=2), it['abv'], ws['spec_size'])
                run(para(body, after=0), it['desc'], ws['desc_size'])
                if it['shot_highball']:
                    sp = para(t.cell(1, 2), first=True,
                              right_tab=price_col - 0.15)
                    run(sp, 'Shot / Highball\t', ws['label_size'])
                    run(sp, it['shot_highball'], ws['price_size'])
                if i < len(sec['items']) - 1:
                    gap = doc.add_paragraph()
                    gf = gap.paragraph_format
                    gf.space_before = gf.space_after = Pt(0)
                    gf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                    gf.line_spacing = Pt(8)

        else:   # etc
            img_col = 1.6
            col_w = (body_w - img_col) / 2
            t = base_table(1, 3)
            cells = t.rows[0].cells
            cells[0].width = Cm(img_col)
            cells[1].width = cells[2].width = Cm(col_w)
            picture(cells[0], sec['img'], 1.17)
            for cell, items in ((cells[1], sec['left']), (cells[2], sec['right'])):
                first = True
                for name, price in items:
                    p = para(cell, first=first, after=10,
                             right_tab=col_w - 0.4)
                    run(p, name + '\t' + price, ws['etc_size'], bold=True)
                    first = False

    tail_paragraph(doc)
    doc.save(os.path.join(DIST, out_name))
    print('saved', out_name)


if __name__ == '__main__':
    build_dish('food', '페피트 요리 메뉴판.docx')
    build_dish('lunch', '페피트 런치 메뉴판.docx')
    build_wine('페피트 와인 메뉴판.docx')
    build_course('course', '페피트 런치 코스메뉴.docx')
    # 섹션 헤더 없이 7코스만 있어 A5 한 장을 채우도록 간격을 넓힌다
    build_course('tasting', '페피트 청담 테이스팅 코스 메뉴판.docx', item_gap=30)
    build_whisky('페피트 위스키 메뉴판.docx')
