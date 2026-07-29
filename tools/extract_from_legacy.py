# -*- coding: utf-8 -*-
"""기존 메뉴판 docx에서 내용을 구조화된 JSON으로 추출."""
import json, re, os
import docx

SRC = r"C:\Users\jj272\OneDrive\문서\카카오톡 받은 파일"
OUT = r"C:\Users\jj272\OneDrive\문서\카카오톡 받은 파일\페피트 메뉴판_새버전\data"
os.makedirs(OUT, exist_ok=True)

PRICE = re.compile(r'([\d,]+)\.\s*$')


def norm(s):
    return re.sub(r'[\t ]+', ' ', s).strip()


def split_price(line):
    """'홍새우                28. Red Shrimp' -> ('홍새우', '28.', 'Red Shrimp')"""
    m = re.search(r'\s([\d,]+)\.\s*(.*)$', line)
    if not m:
        return norm(line), '', ''
    kor = norm(line[:m.start()])
    return kor, m.group(1) + '.', norm(m.group(2))


# ---------------- 요리 메뉴판 ----------------
def extract_food():
    d = docx.Document(os.path.join(SRC, '페피트 요리 메뉴판_A4.docx'))
    lines = [p.text for p in d.paragraphs]
    heads = {'STARTER', 'PASTA', 'MAIN', 'MEAL', 'SIDE', 'DESSERT'}
    sections, cur = [], None
    buf = []

    def flush():
        nonlocal buf
        if not buf or cur is None:
            buf = []
            return
        kor, price, eng = split_price(buf[0])
        rest = buf[1:]
        if not eng and rest and not rest[0].startswith('*'):
            eng = norm(rest.pop(0))
        dk = norm(rest.pop(0).lstrip('* ')) if rest and rest[0].lstrip().startswith('*') else ''
        de = norm(rest.pop(0)) if rest else ''
        cur['items'].append({'kor': kor, 'eng': eng, 'price': price,
                             'desc_kor': dk, 'desc_eng': de})
        buf = []

    for ln in lines:
        t = norm(ln)
        if not t or t == 'PÉPITE':
            continue
        if t in heads:
            flush()
            cur = {'name': t, 'items': []}
            sections.append(cur)
            continue
        if PRICE.search(t) or re.search(r'\s[\d,]+\.\s', t):
            flush()
        buf.append(ln)
    flush()
    return {'title': 'PÉPITE', 'sections': sections}


# ---------------- 런치 메뉴 ----------------
def extract_lunch():
    d = docx.Document(os.path.join(SRC, '페피트 런치메뉴.docx'))
    lines = [p.text for p in d.paragraphs]
    heads = {'STARTER', 'PASTA', 'MAIN', 'MEAL', 'SIDE', 'DESSERT', 'DRINKS'}
    sections, cur, buf = [], None, []

    def flush():
        nonlocal buf
        if not buf or cur is None:
            buf = []
            return
        kor, price, eng = split_price(buf[0])
        rest = buf[1:]
        if not eng and rest and not rest[0].lstrip().startswith('*'):
            eng = norm(rest.pop(0))
        dk = norm(rest.pop(0).lstrip('* ')) if rest and rest[0].lstrip().startswith('*') else ''
        de = norm(rest.pop(0)) if rest else ''
        cur['items'].append({'kor': kor, 'eng': eng, 'price': price,
                             'desc_kor': dk, 'desc_eng': de})
        buf = []

    for ln in lines:
        t = norm(ln)
        if not t or t == 'PÉPITE LUNCH':
            continue
        if t in heads:
            flush()
            cur = {'name': t, 'items': []}
            sections.append(cur)
            continue
        if PRICE.search(t) or re.search(r'\s[\d,]+\.\s', t):
            flush()
        buf.append(ln)
    flush()
    return {'title': 'PÉPITE LUNCH', 'sections': sections}


# ---------------- 와인 메뉴판 ----------------
def extract_wine():
    d = docx.Document(os.path.join(SRC, '페피트 와인 메뉴판 - 존나빡침.docx'))
    lines = [p.text for p in d.paragraphs]
    heads = {'WHITE WINE', 'RED WINE', 'SPARKING WINE', 'SPARKLING WINE'}
    sections, cur, buf = [], None, []

    def flush():
        nonlocal buf
        blk = [b for b in buf if norm(b)]
        buf = []
        if not blk or cur is None:
            return
        kor, price, eng = split_price(blk[0])
        rest = blk[1:]
        if not eng and rest and '>' not in rest[0]:
            eng = norm(rest.pop(0))
        region = norm(rest.pop(0)) if rest and '>' in rest[0] else ''
        dk, de = [], []
        for r in rest:
            (de if re.match(r'^[A-Za-z\u201c"(]', norm(r)) else dk).append(norm(r))
        cur['items'].append({'kor': kor, 'eng': eng, 'price': price,
                             'region': region, 'desc_kor': ' '.join(dk),
                             'desc_eng': ' '.join(de)})

    for ln in lines:
        t = norm(ln)
        if t in heads:
            flush()
            cur = {'name': 'SPARKLING WINE' if t.startswith('SPARK') else t, 'items': []}
            sections.append(cur)
            continue
        if not t:
            flush()
            continue
        if re.search(r'\s[\d,]+\.', t) and ('%' not in t) and '>' not in t:
            flush()
        buf.append(ln)
    flush()
    return {'title': 'PÉPITE WINE', 'sections': sections}


# ---------------- 런치 코스메뉴 ----------------
def extract_course():
    import zipfile
    from lxml import etree

    def W(t):
        return '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}' + t

    root = etree.fromstring(zipfile.ZipFile(
        os.path.join(SRC, '페피트 런치코스메뉴.docx')).read('word/document.xml'))
    tb = [b for b in root.iter(W('txbxContent')) if 'PASTA' in ''.join(b.itertext())][0]
    txt = [norm(''.join(p.itertext())) for p in tb.findall(W('p'))]
    txt = [t for t in txt if t]
    return {'title': 'PÉPITE', 'subtitle': 'LUNCH SET MENU', 'lines': txt}


for name, fn in [('food', extract_food), ('lunch', extract_lunch),
                 ('wine', extract_wine), ('course', extract_course)]:
    data = fn()
    with open(os.path.join(OUT, name + '.json'), 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    if 'sections' in data:
        print(name, [(s['name'], len(s['items'])) for s in data['sections']])
    else:
        print(name, len(data['lines']), '줄')
